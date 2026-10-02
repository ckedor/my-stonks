"""What a UCITS ETF holds, from the file its manager publishes.

No regulator publishes a UCITS fund's holdings, so the manager's own export is
the only source. It is a page made for people, not a filing: nothing promises
its layout. So the reader fails rather than guesses — a header it cannot
find, a date it cannot read, or weights that do not add up to the whole fund
stop the file — because a holdings list read wrong looks exactly like one
read right.

Synchronous and file-bound: run it off the event loop.
"""

from __future__ import annotations

import csv
import io
import re
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.infra.exceptions import IntegrationBadResponse
from app.modules.market_data.domain.etf_registry import is_isin

DWS = 'dws'
ISHARES = 'ishares'

#: Rows searched for the header: the file opens with a title block.
HEADER_SEARCH_ROWS = 30

#: How far the weights may sum from the whole fund. Rounding of a thousand
#: lines and a cash line left out stay inside it; a file cut short does not.
WEIGHT_SUM_TOLERANCE = Decimal('0.05')

_ISO_DATE = re.compile(r'\b(\d{4})-(\d{2})-(\d{2})\b')
#: The UK export writes dates day first.
_DAY_FIRST_DATE = re.compile(r'\b(\d{1,2})[./](\d{1,2})[./](\d{4})\b')
_NAMED_MONTH_DATE = re.compile(r'\b(\d{1,2}) ([A-Za-z]{3,9})\.? (\d{4})\b')


@dataclass
class ManagerHoldingsFile:
    report_date: date
    holdings: list[dict]


def _label(value) -> str:
    """A header cell, compared without case, spacing or punctuation."""
    return re.sub(r'[^a-z]', '', str(value or '').lower())


def _text(value) -> str | None:
    text = str(value).strip() if value is not None else ''
    return text or None


def _number(value) -> tuple[Decimal, bool] | None:
    """The cell's number, and whether it was written as a percentage."""
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, int | float | Decimal):
        return Decimal(str(value)), False
    text = str(value).strip().replace('\u00a0', '').replace(' ', '')
    percent = text.endswith('%')
    text = text.rstrip('%')
    # A lone comma is the decimal mark; next to a point, it groups thousands.
    text = text.replace(',', '.') if ',' in text and '.' not in text else text.replace(',', '')
    try:
        return Decimal(text), percent
    except InvalidOperation:
        return None


def _dates_in(value) -> set[date]:
    if isinstance(value, datetime):
        return {value.date()}
    if isinstance(value, date):
        return {value}
    if not isinstance(value, str):
        return set()
    found: set[date] = set()
    for year, month, day in _ISO_DATE.findall(value):
        found.add(_safe_date(int(year), int(month), int(day)))
    for day, month, year in _DAY_FIRST_DATE.findall(value):
        found.add(_safe_date(int(year), int(month), int(day)))
    for day, month, year in _NAMED_MONTH_DATE.findall(value):
        for pattern in ('%d %b %Y', '%d %B %Y'):
            try:
                found.add(datetime.strptime(f'{day} {month} {year}', pattern).date())
                break
            except ValueError:
                continue
    found.discard(date.min)
    return found


def _safe_date(year: int, month: int, day: int) -> date:
    try:
        return date(year, month, day)
    except ValueError:
        return date.min


#: The header labels read, by role, compared the way `_label` writes them.
#: The first column carrying a role wins.
_ROLES = {
    'name': 'name',
    'isin': 'isin',
    'ticker': 'ticker',
    'issuerticker': 'ticker',
    'country': 'country',
    'currency': 'currency',
    'marketcurrency': 'currency',
}


def _header(rows: list[tuple], required: set[str]) -> tuple[int, dict[str, int]] | None:
    """The header row's index and the columns read, by role."""
    for index, row in enumerate(rows[:HEADER_SEARCH_ROWS]):
        columns: dict[str, int] = {}
        for position, cell in enumerate(row):
            label = _label(cell)
            role = 'weight' if label.startswith('weight') else _ROLES.get(label)
            if role is not None:
                columns.setdefault(role, position)
        if required <= columns.keys():
            return index, columns
    return None


def _no_header(provider: str, rows: list[tuple], expected: str) -> IntegrationBadResponse:
    seen = [[cell for cell in row if cell not in (None, '')] for row in rows[:10]]
    return IntegrationBadResponse(
        f'{provider} holdings file has no {expected} header; first rows: {seen}'[:500],
        provider=provider,
    )


def _cell(row: tuple, position: int | None):
    if position is None or position >= len(row):
        return None
    return row[position]


def _table(rows: list[tuple], header_index: int, columns: dict[str, int], provider: str):
    """The lines under the header, each with its weight as written, and
    whether any weight carried a percent sign.

    The table ends at the first line without a name or a weight: what follows
    is the footnotes. A table cut there by a line that is not a footnote is
    caught by the weights, which then fall short of the fund.
    """
    lines: list[tuple[tuple, Decimal]] = []
    percent_signs = False
    for row in rows[header_index + 1 :]:
        name = _text(_cell(row, columns['name']))
        weight = _number(_cell(row, columns['weight']))
        if name is None or weight is None:
            break
        lines.append((row, weight[0]))
        percent_signs = percent_signs or weight[1]
    if not lines:
        raise IntegrationBadResponse(
            f'{provider} holdings file lists no holdings', provider=provider
        )
    return lines, percent_signs


def _weight_scale(weights: list[Decimal], *, percent_signs: bool, provider: str) -> Decimal:
    """100 when the weights are percentages, 1 when they are ratios.

    Either way they must add up to the whole fund: a file cut short, or a
    column that is not the weight, does not.
    """
    total = sum(weights, Decimal(0))
    if percent_signs or abs(total / 100 - 1) <= WEIGHT_SUM_TOLERANCE:
        scale = Decimal(100)
    elif abs(total - 1) <= WEIGHT_SUM_TOLERANCE:
        scale = Decimal(1)
    else:
        raise IntegrationBadResponse(
            f'{provider} holdings weights add up to {total} over {len(weights)} lines, '
            'neither the whole fund in percent nor as a ratio',
            provider=provider,
        )
    if abs(total / scale - 1) > WEIGHT_SUM_TOLERANCE:
        raise IntegrationBadResponse(
            f'{provider} holdings weights add up to {total}% over {len(weights)} lines',
            provider=provider,
        )
    return scale


def _holding(row: tuple, columns: dict[str, int], weight: Decimal) -> dict:
    isin = (_text(_cell(row, columns.get('isin'))) or '').upper()
    ticker = _text(_cell(row, columns.get('ticker')))
    country = (_text(_cell(row, columns.get('country'))) or '').upper()
    currency = (_text(_cell(row, columns.get('currency'))) or '').upper()
    return {
        'name': (_text(_cell(row, columns['name'])) or '')[:300],
        'isin': isin if is_isin(isin) else None,
        # A dash is how a cash line says it has no ticker.
        'ticker': ticker[:30] if ticker and ticker != '-' else None,
        # The table keeps a code; a country written out is left out rather
        # than translated by guess.
        'country': country if len(country) == 2 and country.isalpha() else None,
        'currency': currency if len(currency) == 3 and currency.isalpha() else None,
        'weight': weight,
    }


def read_dws_constituents(content: bytes) -> ManagerHoldingsFile:
    """DWS's constituents spreadsheet: the date it describes and every line.

    The weight is returned as a ratio, as every proportion in the domain is,
    whether the file writes 5.12, "5.12%" or 0.0512: the file says which by
    its percent sign or, without one, by what the weights add up to.
    """
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError, ValueError) as exc:
        raise IntegrationBadResponse(
            f'DWS constituents file is not an xlsx workbook (starts with {content[:16]!r})',
            provider=DWS,
        ) from exc
    try:
        sheets = [list(sheet.iter_rows(values_only=True)) for sheet in workbook.worksheets]
    finally:
        workbook.close()

    for rows in sheets:
        found = _header(rows, {'name', 'isin', 'weight'})
        if found is not None:
            break
    else:
        raise _no_header(DWS, sheets[0] if sheets else [], 'Name/ISIN/Weighting')
    header_index, columns = found

    dates = {day for row in rows[:header_index] for cell in row for day in _dates_in(cell)}
    if len(dates) != 1:
        raise IntegrationBadResponse(
            f'DWS constituents file states {len(dates)} dates above its header, not one: '
            f'{sorted(dates)}',
            provider=DWS,
        )
    lines, percent_signs = _table(rows, header_index, columns, DWS)
    scale = _weight_scale(
        [weight for _, weight in lines], percent_signs=percent_signs, provider=DWS
    )
    return ManagerHoldingsFile(
        report_date=dates.pop(),
        holdings=[_holding(row, columns, weight / scale) for row, weight in lines],
    )


# --- iShares ------------------------------------------------------------------

#: The line above the header that dates the file: `Fund Holdings as of,"30/Sep/2026"`.
ISHARES_DATE_LABEL = 'fundholdingsasof'

#: How iShares has written that date: the UK site day first with the month
#: named, the US site month first.
_ISHARES_DATE_FORMATS = ('%d/%b/%Y', '%d-%b-%Y', '%d %b %Y', '%b %d, %Y', '%d/%m/%Y', '%Y-%m-%d')


def _ishares_date(value: str | None) -> date | None:
    text = (value or '').strip()
    for pattern in _ISHARES_DATE_FORMATS:
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    return None


def read_ishares_holdings(content: bytes) -> ManagerHoldingsFile:
    """The holdings CSV an iShares product page exports: a block of fund
    facts, a blank line, the table, then the legal text.

    The ISIN is read when the file has the column and is not required: the
    lines are named and weighed without it.
    """
    text = content.decode('utf-8-sig', errors='replace')
    if text.lstrip()[:1] == '<':
        raise IntegrationBadResponse(
            f'iShares answered a page, not the holdings CSV: {text.strip()[:80]!r}',
            provider=ISHARES,
        )
    rows = [tuple(row) for row in csv.reader(io.StringIO(text))]

    found = _header(rows, {'name', 'weight'})
    if found is None:
        raise _no_header(ISHARES, rows, 'Name/Weight (%)')
    header_index, columns = found

    dates = {
        _ishares_date(row[1] if len(row) > 1 else None)
        for row in rows[:header_index]
        if row and _label(row[0]) == ISHARES_DATE_LABEL
    }
    if len(dates) != 1 or None in dates:
        raise IntegrationBadResponse(
            'iShares holdings file has no readable "Fund Holdings as of" line above its header',
            provider=ISHARES,
        )
    lines, percent_signs = _table(rows, header_index, columns, ISHARES)
    scale = _weight_scale(
        [weight for _, weight in lines], percent_signs=percent_signs, provider=ISHARES
    )
    return ManagerHoldingsFile(
        report_date=dates.pop(),
        holdings=[_holding(row, columns, weight / scale) for row, weight in lines],
    )
